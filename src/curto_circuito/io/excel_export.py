"""Exportação de resultados para Excel (openpyxl)."""

from __future__ import annotations
import cmath
import math
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from ..models.network import Network
from ..models.results import StudyResults
from ..utils.units import z_to_pu

_HEADER_FILL = PatternFill("solid", fgColor="1F497D")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_MT_FILL = PatternFill("solid", fgColor="D6E4F0")
_BT_FILL = PatternFill("solid", fgColor="D5F5E3")
_MT_THRESHOLD_KV = 1.0


def _header_row(ws, values: list[str]) -> None:
    ws.append(values)
    for cell in ws[ws.max_row]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", wrap_text=True)


def _fmt(val) -> str | float:
    if val is None:
        return "-"
    if val == float("inf"):
        return "∞"
    return round(val, 4)


def export(results: StudyResults, network: Network, path: str | Path) -> None:
    wb = Workbook()

    # ---- Aba Resultados ----
    ws_res = wb.active
    ws_res.title = "Resultados"
    ws_res.row_dimensions[1].height = 30
    _header_row(ws_res, [
        "Nível", "Barra", "Un (kV)",
        "Icc3φ (kA)", "Icc pico (kA)", "Icc3φ-g (kA)",
        "Icc2φ (kA)", "Icc1φ (kA)", "Icc2φ-g (kA)", "κ",
        "θ3φ (°)", "Desl. (°)",
    ])

    buses_mt = [b for b in results.buses if b.un_kv > _MT_THRESHOLD_KV]
    buses_bt = [b for b in results.buses if b.un_kv <= _MT_THRESHOLD_KV]
    for br in buses_mt + buses_bt:
        is_mt = br.un_kv > _MT_THRESHOLD_KV
        fill = _MT_FILL if is_mt else _BT_FILL
        f3  = br.get_fault("3F")
        f3t = br.get_fault("3F-T")
        f2  = br.get_fault("2F")
        f1  = br.get_fault("1F-T")
        f2e = br.get_fault("2F-T")
        ws_res.append([
            "MT" if is_mt else "BT",
            br.node_name,
            round(br.un_kv, 2),
            _fmt(f3.ik_pp_ka   if f3  else None),
            _fmt(f3.ip_ka      if f3  else None),
            _fmt(f3t.ik_pp_ka  if f3t else None),
            _fmt(f2.ik_pp_ka   if f2  else None),
            _fmt(f1.ik_pp_ka   if f1  else None),
            _fmt(f2e.ik_pp_ka  if f2e else None),
            _fmt(f3.kappa      if f3  else None),
            _fmt(f3.angle_deg  if f3  else None),
            round(br.phase_shift_deg, 1),
        ])
        for cell in ws_res[ws_res.max_row]:
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center")

    # ---- Aba Impedâncias ----
    ws_imp = wb.create_sheet("Impedâncias")
    _header_row(ws_imp, [
        "Nível", "Barra", "Un (kV)",
        "R1 (mΩ)", "X1 (mΩ)", "|Z1| (mΩ)", "Z1_pu", "arg(Z1) (°)",
        "R0 (mΩ)", "X0 (mΩ)", "|Z0| (mΩ)", "Z0_pu", "arg(Z0) (°)",
    ])
    for br in results.buses:
        f3 = br.get_fault("3F")
        f1 = br.get_fault("1F-T")
        if f3:
            z1 = f3.z1_ohm * 1000
            z0 = (f1.z0_ohm if f1 else 0+0j) * 1000
            z1_pu = z_to_pu(f3.z1_ohm, br.un_kv, results.s_base_mva)
            z0_pu = z_to_pu(f3.z0_ohm, br.un_kv, results.s_base_mva)
            is_mt = br.un_kv > _MT_THRESHOLD_KV
            ws_imp.append([
                "MT" if is_mt else "BT",
                br.node_name, round(br.un_kv, 2),
                round(z1.real, 4), round(z1.imag, 4), round(abs(z1), 4),
                round(abs(z1_pu), 6), round(math.degrees(cmath.phase(z1_pu)), 2),
                round(z0.real, 4), round(z0.imag, 4), round(abs(z0), 4),
                round(abs(z0_pu), 6), round(math.degrees(cmath.phase(z0_pu)), 2),
            ])
            fill = _MT_FILL if is_mt else _BT_FILL
            for cell in ws_imp[ws_imp.max_row]:
                cell.fill = fill
                cell.alignment = Alignment(horizontal="center")

    # ---- Aba p.u. ----
    ws_pu = wb.create_sheet("p.u.")
    _header_row(ws_pu, [
        "Nível", "Barra", "Un (kV)",
        "|Z1| (p.u.)", "|Z0| (p.u.)",
        "θ3φ (°)", "θ2φ (°)", "θ1φ (°)", "θ2φ-g (°)",
        "Desl. (°)",
    ])
    for br in buses_mt + buses_bt:
        is_mt = br.un_kv > _MT_THRESHOLD_KV
        f3  = br.get_fault("3F")
        f2  = br.get_fault("2F")
        f1  = br.get_fault("1F-T")
        f2e = br.get_fault("2F-T")
        z1_pu = abs(z_to_pu(f3.z1_ohm, br.un_kv, results.s_base_mva)) if f3 else None
        z0_pu = abs(z_to_pu(f3.z0_ohm, br.un_kv, results.s_base_mva)) if f3 else None
        ws_pu.append([
            "MT" if is_mt else "BT",
            br.node_name,
            round(br.un_kv, 2),
            _fmt(z1_pu),
            _fmt(z0_pu),
            _fmt(f3.angle_deg  if f3  else None),
            _fmt(f2.angle_deg  if f2  else None),
            _fmt(f1.angle_deg  if f1  else None),
            _fmt(f2e.angle_deg if f2e else None),
            round(br.phase_shift_deg, 1),
        ])
        fill = _MT_FILL if is_mt else _BT_FILL
        for cell in ws_pu[ws_pu.max_row]:
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center")

    # ---- Aba Rede ----
    ws_net = wb.create_sheet("Rede")
    _header_row(ws_net, ["Tipo", "ID", "Nome", "Parâmetros"])
    for node in network.nodes.values():
        comp = node.component
        if comp is not None:
            ws_net.append([type(comp).__name__, comp.id, comp.name, str(comp.__dict__)])

    # Auto-width
    for ws in [ws_res, ws_imp, ws_pu, ws_net]:
        for col in ws.columns:
            max_len = max((len(str(c.value or "")) for c in col), default=10)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 2, 40)

    wb.save(path)
